"""Compare completed action receipts and their saved numerical/display data.

Only explicitly supplied output roots and declared creation/runtime metadata
are normalized. Scientific values/settings and engine fingerprints stay exact.
SVG files retain native renderer validation; their source tables are compared
because provenance and vector identifiers are intentionally run-specific.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import re
import statistics
from benchmark_pipeline_execution import digest


def tables(receipt):
    data=json.loads(receipt.read_text())
    return data,{name:receipt.parent/name for name in data.get('outputs',{})
                 if Path(name).suffix.lower() in {'.csv','.tif','.tiff','.npy','.npz','.png','.jpg','.jpeg','.mp4','.xlsx','.html','.zip'}
                 or Path(name).name in {'rhythm.json','instrumental_control.json','eligibility.json'}}


def equivalent(left,right,left_root,right_root,metadata_aliases=None):
    if left.suffix.lower()=='.zip':
        import zipfile
        with zipfile.ZipFile(left) as a,zipfile.ZipFile(right) as b:
            names_a,names_b=a.namelist(),b.namelist()
            valid=bool(names_a) and len(names_a)==len(set(names_a)) and len(names_b)==len(set(names_b))
            attributes=('comment','extra','external_attr','internal_attr','flag_bits','compress_type','create_system')
            same=valid and names_a==names_b and a.comment==b.comment and all(
                a.read(n)==b.read(n) and all(getattr(a.getinfo(n),key)==getattr(b.getinfo(n),key)
                                           for key in attributes) for n in names_a)
            return same,'exact ordered archive members and bytes; ZIP creation timestamps ignored'
    if digest(left)==digest(right):return True,'byte-identical'
    import numpy as np
    suffix=left.suffix.lower()
    if suffix in {'.csv','.html','.json'}:
        a,b=left.read_text(encoding='utf-8'),right.read_text(encoding='utf-8')
        for root in (left_root,right_root):
            for spelling in (str(root),root.as_posix(),str(root).replace('\\','\\\\')):
                a=a.replace(spelling,'<trial>');b=b.replace(spelling,'<trial>')
        if metadata_aliases and (left.name=='entries.csv' or left.name.endswith('_display.csv')):
            import csv,io
            x,y=csv.DictReader(io.StringIO(a)),csv.DictReader(io.StringIO(b))
            if x.fieldnames!=y.fieldnames:return False,'different display columns'
            rows_x,rows_y=list(x),list(y)
            for rows,aliases in zip((rows_x,rows_y),metadata_aliases):
                for row in rows:
                    for column,values in aliases.items():
                        value=row.get(column)
                        if value in values:row[column]=values[value]
            return rows_x==rows_y,'exact display data; verified report references and confirmation creation time normalized'
        volatile=None
        if left.name=='_runs_index.csv':
            volatile='seconds'  # The runtime being compared, not scientific time.
        elif left.name=='sources.csv' and left.parent.name=='data' and left.parent.parent.name.endswith('_bundle'):
            volatile='modified'  # Preserve source byte count and content digest.
        if volatile:
            import csv,io
            x,y=csv.DictReader(io.StringIO(a)),csv.DictReader(io.StringIO(b))
            if x.fieldnames!=y.fieldnames:return False,'different metadata table columns'
            rows_x,rows_y=list(x),list(y)
            for row in rows_x+rows_y:row.pop(volatile,None)
            return rows_x==rows_y,'exact metadata after trial paths and declared volatile '+volatile+' field'
        if left.name.endswith('_statistics.csv'):
            import csv,io
            from datetime import datetime
            csv.field_size_limit(20_000_000)
            x,y=csv.DictReader(io.StringIO(a)),csv.DictReader(io.StringIO(b))
            if x.fieldnames!=y.fieldnames:return False,'different statistics columns'
            if 'results_json' in (x.fieldnames or []):
                rows_x,rows_y=list(x),list(y)
                for row in rows_x+rows_y:
                    if not row['results_json']:continue
                    records=json.loads(row['results_json'])
                    if isinstance(records,list):
                        for record in records:
                            provenance=record.get('provenance',{}) if isinstance(record,dict) else {}
                            created=provenance.get('generated_utc')
                            if isinstance(created,str):
                                datetime.fromisoformat(created)  # Refuse malformed metadata.
                                provenance['generated_utc']='<generation-time>'
                    row['results_json']=json.dumps(records,sort_keys=True)
                return rows_x==rows_y,'exact statistics and records; only provenance.generated_utc creation time normalized'
        return a==b,'exact text after explicit trial-path normalization'
    if suffix in {'.tif','.tiff'}:
        import tifffile
        a,b=tifffile.imread(left),tifffile.imread(right)
        return a.dtype==b.dtype and np.array_equal(a,b,equal_nan=True),'exact typed pixel arrays'
    if suffix=='.npy':
        a,b=np.load(left,allow_pickle=False),np.load(right,allow_pickle=False)
        return a.dtype==b.dtype and np.array_equal(a,b,equal_nan=True),'exact typed arrays'
    if suffix=='.npz':
        with np.load(left,allow_pickle=False) as a,np.load(right,allow_pickle=False) as b:
            return set(a.files)==set(b.files) and all(a[k].dtype==b[k].dtype and np.array_equal(a[k],b[k],equal_nan=True) for k in a.files),'exact named arrays'
    if suffix in {'.png','.jpg','.jpeg'}:
        from PIL import Image
        with Image.open(left) as a,Image.open(right) as b:
            return a.mode==b.mode and np.array_equal(np.asarray(a),np.asarray(b)),'exact decoded pixels'
    if suffix=='.mp4':
        import imageio.v2 as imageio
        from itertools import zip_longest
        with imageio.get_reader(str(left),'ffmpeg') as a,imageio.get_reader(str(right),'ffmpeg') as b:
            same=a.get_meta_data().get('fps')==b.get_meta_data().get('fps')
            for x,y in zip_longest(a,b):
                if x is None or y is None or not np.array_equal(x,y):return False,'different decoded video frames'
            return same,'exact decoded frames and frame rate'
    if suffix=='.xlsx':
        import openpyxl
        a=openpyxl.load_workbook(left,read_only=True,data_only=False)
        b=openpyxl.load_workbook(right,read_only=True,data_only=False)
        try:
            same=a.sheetnames==b.sheetnames
            for name in a.sheetnames:
                if name not in b:return False,'missing workbook sheet'
                def cells(sheet):
                    return [[(c.value,c.data_type,c.number_format) for c in row] for row in sheet.iter_rows()]
                same=same and cells(a[name])==cells(b[name])
            return same,'exact sheet order, cells, formulas, types and number formats'
        finally:a.close();b.close()
    return False,'different binary output; requires inspection'


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('before',type=Path);parser.add_argument('after',type=Path);parser.add_argument('output',type=Path)
    args=parser.parse_args();a,af=tables(args.before);b,bf=tables(args.after)
    checks=[]
    for key in sorted(set(af)|set(bf)):
        if key not in af or key not in bf:checks.append({'file':key,'equal':False,'comparison':'missing counterpart'});continue
        if digest(af[key])!=a['outputs'][key]['sha256'] or digest(bf[key])!=b['outputs'][key]['sha256']:
            checks.append({'file':key,'equal':False,'comparison':'output changed since its execution receipt'});continue
        equal,method=equivalent(af[key],bf[key],args.before.parent.parent.parent,args.after.parent.parent.parent)
        checks.append({'file':key,'equal':bool(equal),'comparison':method})
    accepted=a.get('status')==b.get('status')=='measured' and a.get('id')==b.get('id')
    environment=a.get('ledger_implementation',{}).get('dependency_environment')
    stable=bool(environment) and environment==b.get('ledger_implementation',{}).get('dependency_environment')
    stable=stable and all(row.get('source_unchanged') and row.get('dependencies_unchanged') for row in (a,b))
    same_sampling=a.get('memory_sampling')==b.get('memory_sampling')
    stable=stable and same_sampling
    source_differences=sorted(k for k in set(a.get('source_before',{}))|set(b.get('source_before',{}))
        if a.get('source_before',{}).get(k)!=b.get('source_before',{}).get(k))
    # These two entry points are the intentional provenance-batching treatment;
    # their actual loaded hashes are recorded separately by configure_ledger.
    treatments={'src/pymicroglia/run.py':'dispatcher_sha256',
                'src/pymicroglia/measure/__init__.py':'measurement_entry_sha256'}
    stable=stable and set(source_differences)<=set(treatments)
    if source_differences:
        stable=stable and a.get('ledger_implementation',{}).get('mode')=='baseline' and b.get('ledger_implementation',{}).get('mode')=='current'
        stable=stable and all(len(row.get('ledger_implementation',{}).get('dispatcher_sha256',''))==64 for row in (a,b))
        for source in source_differences:
            if source in treatments:
                stable=stable and all(len(row.get('ledger_implementation',{}).get(
                    treatments[source],row.get('source_before',{}).get(source,'')))==64 for row in (a,b))
    def settings(row,root):
        value=json.dumps({k:row[k] for k in ('params','overrides','fixture') if k in row},sort_keys=True)
        for spelling in (str(root),root.as_posix()):
            value=value.replace(json.dumps(spelling)[1:-1],'<trial>')
        return value
    same_settings=settings(a,args.before.parent.parent.parent)==settings(b,args.after.parent.parent.parent)
    inputs=[]
    if a.get('input_snapshot_sha256') or b.get('input_snapshot_sha256'):
        inputs.append({'source':'saved fixture snapshot','equal':a.get('input_snapshot_sha256')==b.get('input_snapshot_sha256')})
    left_inputs=args.before.parent.parent/'inputs'
    right_inputs=args.after.parent.parent/'inputs'
    if left_inputs.is_dir() and right_inputs.is_dir():
        left={p.relative_to(left_inputs).as_posix():p for p in left_inputs.rglob('*') if p.is_file()}
        right={p.relative_to(right_inputs).as_posix():p for p in right_inputs.rglob('*') if p.is_file()}
        for key in sorted(set(left)|set(right)):
            equal,method=(False,'missing input') if key not in left or key not in right else equivalent(
                left[key],right[key],args.before.parent.parent.parent,args.after.parent.parent.parent)
            inputs.append({'source':key,'equal':bool(equal),'comparison':method})
    result={'id':a.get('id'),'kind':'execution','status':'measured' if accepted else 'failed',
        'before_seconds':[a.get('seconds')],'after_seconds':[b.get('seconds')],
        'peak_rss_bytes':max(a.get('peak_rss_bytes',0),b.get('peak_rss_bytes',0)),
        'before_peak_rss_bytes':a.get('peak_rss_bytes'),'after_peak_rss_bytes':b.get('peak_rss_bytes'),
        'matched_stable_environment':bool(stable),
        'application_source_differences':source_differences,
        'same_settings':same_settings,
        'same_memory_sampling':same_sampling,
        'input_checks':inputs,
        'parity':accepted and stable and same_settings and all(c['equal'] for c in inputs) and bool(checks) and all(c['equal'] for c in checks),
        'checks':checks,'before_receipt':str(args.before),'after_receipt':str(args.after),
        'before_receipt_sha256':digest(args.before),'after_receipt_sha256':digest(args.after)}
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k!='checks'}))
    return 0 if result['parity'] else 1


if __name__=='__main__':raise SystemExit(main())
